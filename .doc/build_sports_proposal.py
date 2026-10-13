from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = Path(r"E:\Store web\Ho-so-y-tuong-website-the-thao.docx")

BLACK = "171717"
ORANGE = "F4511E"
WHITE = "FFFFFF"
LIGHT = "F7F7F7"
BORDER = "D9D9D9"
MUTED = "737373"
PALE_ORANGE = "FFF1EC"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Arial", size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def remove_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def remove_style_border(style):
    p_pr = style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=8.5, color=MUTED)


def add_kicker(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text.upper())
    set_run_font(run, size=9, bold=True, color=ORANGE)
    keep_with_next(p)
    return p


def add_title(doc, text, size=24):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True, color=BLACK)
    keep_with_next(p)
    remove_paragraph_border(p)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_run_font(run, size=15 if level == 1 else 12, bold=True, color=BLACK)
    keep_with_next(p)
    return p


def add_body(doc, text, bold_lead=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.18
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        set_run_font(first, size=10.5, bold=True, color=BLACK)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, size=10.5, color=BLACK)
    else:
        run = p.add_run(text)
        set_run_font(run, size=10.5, color=BLACK)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Cm(0.55 + 0.45 * level)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    set_run_font(run, size=10.2, color=BLACK)
    return p


def add_numbered(doc, number, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    r1 = p.add_run(f"{number}  {title}  ")
    set_run_font(r1, size=10.3, bold=True, color=ORANGE)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.3, color=BLACK)
    return p


def format_table(table, header=True, widths=None, font_size=9.2):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        if row_index == 0 and header:
            set_repeat_table_header(row)
        for col_index, cell in enumerate(row.cells):
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths and col_index < len(widths):
                cell.width = Cm(widths[col_index])
            if header and row_index == 0:
                set_cell_shading(cell, BLACK)
            elif row_index % 2 == 0:
                set_cell_shading(cell, LIGHT)
            else:
                set_cell_shading(cell, WHITE)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                for run in p.runs:
                    set_run_font(
                        run,
                        size=font_size,
                        bold=(header and row_index == 0),
                        color=WHITE if header and row_index == 0 else BLACK,
                    )


def add_section_intro(doc, page_no, title, intro):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.05
    kicker = p.add_run(f"{page_no:02d}  ")
    set_run_font(kicker, size=10, bold=True, color=ORANGE)
    title_run = p.add_run(title)
    set_run_font(title_run, size=22, bold=True, color=BLACK)
    keep_with_next(p)
    remove_paragraph_border(p)
    add_body(doc, intro, space_after=10)


def add_page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.55)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)
section.header_distance = Cm(0.7)
section.footer_distance = Cm(0.65)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(BLACK)

for style_name in ("Title", "Heading 1", "Heading 2"):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.color.rgb = RGBColor.from_string(BLACK)
    remove_style_border(style)

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer_p.paragraph_format.space_before = Pt(0)
footer_p.paragraph_format.space_after = Pt(0)
footer_run = footer_p.add_run("Hồ sơ ý tưởng website thể thao   |   ")
set_run_font(footer_run, size=8.5, color=MUTED)
add_page_field(footer_p)

# Page 1 Cover
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(74)
r = p.add_run("ĐỀ XUẤT Ý TƯỞNG KINH DOANH SỐ")
set_run_font(r, size=9, bold=True, color=ORANGE)

p = doc.add_paragraph(style="Title")
p.paragraph_format.space_after = Pt(16)
p.paragraph_format.line_spacing = 1.0
r = p.add_run("Website kinh doanh\ndụng cụ thể thao")
set_run_font(r, size=34, bold=True, color=BLACK)
remove_paragraph_border(p)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(44)
p.paragraph_format.line_spacing = 1.25
r = p.add_run("Hồ sơ định hướng sản phẩm, trải nghiệm mua sắm và lộ trình triển khai")
set_run_font(r, size=14, color=MUTED)

meta = doc.add_table(rows=3, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.LEFT
meta.autofit = False
meta_data = [
    ("Thương hiệu", "[TÊN THƯƠNG HIỆU]"),
    ("Khách hàng", "[TÊN KHÁCH HÀNG]"),
    ("Liên hệ", "[THÔNG TIN LIÊN HỆ]"),
]
for idx, (label, value) in enumerate(meta_data):
    meta.cell(idx, 0).width = Cm(4.0)
    meta.cell(idx, 1).width = Cm(10.5)
    for cell in meta.rows[idx].cells:
        set_cell_margins(cell, top=100, start=0, bottom=100, end=120)
        set_cell_border(cell, color=WHITE, size="0")
    p0 = meta.cell(idx, 0).paragraphs[0]
    p0.clear()
    rr = p0.add_run(label.upper())
    set_run_font(rr, size=8.5, bold=True, color=MUTED)
    p1 = meta.cell(idx, 1).paragraphs[0]
    p1.clear()
    rr = p1.add_run(value)
    set_run_font(rr, size=10.5, bold=True, color=BLACK)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(66)
r = p.add_run("Bản đề xuất có thể tùy chỉnh cho từng khách hàng")
set_run_font(r, size=9.5, italic=True, color=MUTED)

add_page_break(doc)

# Page 2 Overview and contents
add_section_intro(
    doc,
    2,
    "Mục tiêu của đề xuất",
    "Tài liệu này giúp khách hàng hình dung một cửa hàng thể thao trực tuyến có thể bán hàng, xây niềm tin và mở rộng dịch vụ theo từng giai đoạn. Hướng đề xuất ưu tiên trải nghiệm chọn sản phẩm rõ ràng, hình ảnh sản phẩm nổi bật và quy trình mua ngắn gọn trên cả máy tính lẫn điện thoại.",
)
add_heading(doc, "Kết luận đề xuất", 1)
add_body(
    doc,
    "Chọn hướng Performance Store làm nền tảng: giao diện gọn, mạnh, tập trung vào thông số và lợi ích của sản phẩm. Bổ sung nội dung cộng đồng ở mức vừa đủ để hỗ trợ câu lạc bộ, đội nhóm và giữ quan hệ với khách hàng sau mua.",
)
add_heading(doc, "Mục lục", 1)
toc_rows = [
    ("03", "Khách hàng mục tiêu và định vị"),
    ("04", "Ba hướng ý tưởng thương hiệu"),
    ("05", "Ý tưởng kinh doanh nổi bật"),
    ("06", "Hành trình mua hàng"),
    ("07", "Cấu trúc website"),
    ("08", "Kiến trúc trang chủ"),
    ("09", "Trải nghiệm và tương tác"),
    ("10", "Hệ thống hình ảnh và nội dung"),
    ("11", "Lộ trình triển khai"),
    ("12", "Mô hình dữ liệu và bước tiếp theo"),
]
table = doc.add_table(rows=len(toc_rows), cols=2)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
table.autofit = False
for i, (page, label) in enumerate(toc_rows):
    table.cell(i, 0).width = Cm(1.2)
    table.cell(i, 1).width = Cm(14.2)
    for cell in table.rows[i].cells:
        set_cell_border(cell, color=WHITE, size="0")
        set_cell_margins(cell, top=60, start=0, bottom=60, end=100)
    p0 = table.cell(i, 0).paragraphs[0]
    p0.clear()
    rr = p0.add_run(page)
    set_run_font(rr, size=9.5, bold=True, color=ORANGE)
    p1 = table.cell(i, 1).paragraphs[0]
    p1.clear()
    rr = p1.add_run(label)
    set_run_font(rr, size=10, color=BLACK)

add_page_break(doc)

# Page 3 Audience and positioning
add_section_intro(
    doc,
    3,
    "Khách hàng mục tiêu và định vị",
    "Website cần phục vụ cả người đã biết rõ món đồ mình muốn mua và người còn cần tư vấn. Cấu trúc nội dung vì vậy phải cho phép tìm nhanh theo tên sản phẩm, đồng thời dẫn dắt theo môn chơi, trình độ và ngân sách.",
)
audiences = [
    ("Người chơi phong trào", "Muốn chọn đúng sản phẩm, dễ so sánh và có mức giá phù hợp."),
    ("Người chơi bán chuyên", "Quan tâm thông số, độ phù hợp với lối chơi và sản phẩm mới."),
    ("Phụ huynh", "Cần hướng dẫn rõ, lựa chọn an toàn và mức đầu tư hợp lý cho trẻ."),
    ("Câu lạc bộ và đội nhóm", "Cần mua theo bộ, đặt số lượng lớn và nhận tư vấn thống nhất."),
    ("Khách mua quà", "Cần gợi ý nhanh theo người nhận, dịp tặng và khoảng ngân sách."),
]
table = doc.add_table(rows=1, cols=2)
table.rows[0].cells[0].text = "Nhóm khách hàng"
table.rows[0].cells[1].text = "Nhu cầu chính"
for a, need in audiences:
    cells = table.add_row().cells
    cells[0].text = a
    cells[1].text = need
format_table(table, widths=[5.0, 11.0], font_size=9.4)

add_heading(doc, "Định vị đề xuất", 1)
add_body(
    doc,
    "[TÊN THƯƠNG HIỆU] là cửa hàng thể thao đa môn giúp khách chọn sản phẩm chính hãng theo đúng nhu cầu sử dụng, với thông tin dễ hiểu và quy trình mua nhanh.",
)
add_heading(doc, "Ba nguyên tắc vận hành", 2)
add_bullet(doc, "Giải thích sản phẩm theo ngôn ngữ của người chơi, không chỉ liệt kê thông số kỹ thuật.")
add_bullet(doc, "Cho khách thấy bằng chứng về nguồn gốc, chính sách đổi trả và dịch vụ sau mua.")
add_bullet(doc, "Tạo nhiều đường mua: tìm kiếm trực tiếp, chọn theo môn chơi, chọn theo trình độ hoặc mua theo combo.")

add_page_break(doc)

# Page 4 Concept directions
add_section_intro(
    doc,
    4,
    "Ba hướng ý tưởng thương hiệu",
    "Ba hướng dưới đây dùng chung nền tảng thương mại điện tử nhưng khác cách kể câu chuyện và ưu tiên nội dung. Khách hàng có thể chọn một hướng chính trước khi phát triển giao diện chi tiết.",
)
concepts = [
    ("Performance Store", "Hiện đại và mạnh", "Thông số, hiệu suất, so sánh sản phẩm, bộ sưu tập mới", "Phù hợp nhất"),
    ("Sport Lifestyle", "Năng động và thời trang", "Trang phục, giày, phối đồ, hình ảnh sử dụng thực tế", "Phù hợp khi thời trang chiếm tỷ trọng cao"),
    ("Club Commerce", "Cộng đồng và gắn kết", "Câu lạc bộ, đội nhóm, đặt số lượng lớn, hoạt động cộng đồng", "Phù hợp để mở rộng sau"),
]
table = doc.add_table(rows=1, cols=4)
for i, value in enumerate(("Hướng", "Cảm nhận", "Nội dung ưu tiên", "Vai trò")):
    table.rows[0].cells[i].text = value
for row in concepts:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
format_table(table, widths=[3.5, 3.6, 6.1, 3.2], font_size=8.8)

add_heading(doc, "Hướng khuyến nghị", 1)
add_body(
    doc,
    "Performance Store nên là lớp nhận diện chính vì giúp sản phẩm trở thành trọng tâm và phù hợp với nhiều nhóm hàng. Club Commerce được dùng như lớp mở rộng cho dịch vụ đội nhóm, sự kiện và khách hàng thân thiết.",
)
add_heading(doc, "Dấu hiệu nhận biết", 2)
add_numbered(doc, "01", "Sản phẩm rõ", "Ảnh lớn, nền sạch, thông số quan trọng xuất hiện đúng lúc.")
add_numbered(doc, "02", "Hành động rõ", "Mỗi khu vực có một mục tiêu chính như xem sản phẩm, so sánh hoặc thêm vào giỏ.")
add_numbered(doc, "03", "Nhịp điệu thể thao", "Typography chắc, khoảng trắng rộng và chuyển động ngắn khi tương tác.")

add_page_break(doc)

# Page 5 Business ideas
add_section_intro(
    doc,
    5,
    "Ý tưởng kinh doanh nổi bật",
    "Các ý tưởng được sắp theo tác động đối với quyết định mua và khả năng triển khai. Giai đoạn đầu nên tập trung vào những tính năng giúp khách chọn đúng sản phẩm và tin tưởng cửa hàng.",
)
ideas = [
    ("Bộ chọn sản phẩm", "Chọn theo môn, trình độ, ngân sách và lối chơi.", "Giảm bối rối khi danh mục lớn."),
    ("Combo sẵn sàng ra sân", "Gộp vợt, giày, túi và phụ kiện tương thích.", "Tăng giá trị đơn và giúp người mới mua nhanh."),
    ("Xác thực chính hãng", "Nhãn xác thực, nguồn gốc và hướng dẫn kiểm tra.", "Tăng niềm tin trước khi thanh toán."),
    ("Trang theo môn chơi", "Không gian riêng cho cầu lông, pickleball và tennis.", "Giúp chiến dịch quảng cáo đi đúng nhu cầu."),
    ("Đơn hàng đội nhóm", "Yêu cầu báo giá, đồng phục và mua số lượng lớn.", "Mở thêm nguồn doanh thu B2B."),
    ("Ra mắt giới hạn", "Đăng ký nhận thông báo và đặt trước sản phẩm mới.", "Tạo nhu cầu mà không cần giảm giá liên tục."),
]
table = doc.add_table(rows=1, cols=3)
for i, value in enumerate(("Ý tưởng", "Cách hoạt động", "Giá trị kinh doanh")):
    table.rows[0].cells[i].text = value
for row in ideas:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
format_table(table, widths=[4.1, 6.1, 6.1], font_size=8.7)

add_heading(doc, "Khả năng mở rộng", 1)
add_body(
    doc,
    "Sau khi luồng bán hàng ổn định, website có thể bổ sung thành viên, điểm thưởng, ưu đãi sinh nhật, đặt lịch căng vợt, bảo hành điện tử và quản lý nhiều chi nhánh. Nội dung tư vấn tiếp tục đóng vai trò kéo khách quay lại mà không phụ thuộc hoàn toàn vào khuyến mãi.",
)

add_page_break(doc)

# Page 6 Journey
add_section_intro(
    doc,
    6,
    "Hành trình mua hàng",
    "Hành trình được thiết kế để khách luôn có bước tiếp theo rõ ràng. Người mua có thể đi thẳng từ tìm kiếm đến sản phẩm hoặc bắt đầu bằng môn chơi và nhận gợi ý phù hợp.",
)
journey = [
    ("1", "Khám phá", "Tìm kiếm, quảng cáo hoặc nội dung tư vấn", "Thấy đúng môn chơi và sản phẩm đáng chú ý"),
    ("2", "Thu hẹp", "Bộ lọc, bộ chọn sản phẩm, danh mục theo môn", "Có danh sách phù hợp với nhu cầu"),
    ("3", "Đánh giá", "Ảnh, thông số, đánh giá, so sánh, xác thực", "Hiểu vì sao sản phẩm phù hợp"),
    ("4", "Quyết định", "Giá, ưu đãi, tồn kho, giao hàng, đổi trả", "Đủ thông tin để thêm vào giỏ"),
    ("5", "Hoàn tất", "Giỏ hàng và quy trình đặt mua ngắn", "Đặt hàng với ít thao tác"),
    ("6", "Sau mua", "Hướng dẫn sử dụng, bảo hành, thành viên", "Quay lại mua phụ kiện hoặc sản phẩm tiếp theo"),
]
table = doc.add_table(rows=1, cols=4)
for i, value in enumerate(("STT", "Giai đoạn", "Điểm chạm", "Kết quả mong muốn")):
    table.rows[0].cells[i].text = value
for row in journey:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
format_table(table, widths=[1.5, 3.0, 5.9, 5.9], font_size=8.6)

add_heading(doc, "Nguyên tắc giảm ma sát", 1)
add_bullet(doc, "Giữ trạng thái bộ lọc khi khách quay lại danh sách sản phẩm.")
add_bullet(doc, "Hiển thị phí giao hàng, đổi trả và tình trạng tồn kho trước bước thanh toán.")
add_bullet(doc, "Cho phép lưu sản phẩm yêu thích mà không bắt buộc đăng nhập ngay.")
add_bullet(doc, "Giữ nút thêm vào giỏ dễ thấy trên điện thoại nhưng không che nội dung quan trọng.")

add_page_break(doc)

# Page 7 Site structure
add_section_intro(
    doc,
    7,
    "Cấu trúc website",
    "Kiến trúc thông tin chia website theo mục đích của khách thay vì chỉ theo cơ cấu nội bộ. Các trang bán hàng, tư vấn và dịch vụ liên kết với nhau bằng ngữ cảnh sản phẩm.",
)
structure = [
    ("Trang chủ", "Điểm vào chính", "Thương hiệu, tìm kiếm, danh mục, sản phẩm và chiến dịch"),
    ("Theo môn chơi", "Khám phá có định hướng", "Cầu lông, pickleball, tennis và bộ sưu tập liên quan"),
    ("Danh mục", "Thu hẹp lựa chọn", "Vợt, giày, quần áo, túi và phụ kiện"),
    ("Chi tiết sản phẩm", "Đánh giá và quyết định", "Ảnh, biến thể, thông số, đánh giá, giao hàng và sản phẩm liên quan"),
    ("Khuyến mãi", "Kích hoạt nhu cầu", "Flash sale, hàng mới, bán chạy và bộ sưu tập theo mùa"),
    ("Tư vấn", "Giáo dục và tìm kiếm tự nhiên", "Hướng dẫn chọn, sử dụng, bảo quản và phòng tránh chấn thương"),
    ("Dịch vụ", "Mở rộng doanh thu", "Đội nhóm, câu lạc bộ, căng vợt và bảo hành"),
    ("Tài khoản", "Quan hệ sau mua", "Đơn hàng, yêu thích, thành viên và thông tin bảo hành"),
]
table = doc.add_table(rows=1, cols=3)
for i, value in enumerate(("Nhóm trang", "Vai trò", "Nội dung chính")):
    table.rows[0].cells[i].text = value
for row in structure:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
format_table(table, widths=[3.6, 4.2, 8.5], font_size=8.4)

add_heading(doc, "Điều hướng chính", 1)
add_body(doc, "Menu cấp cao gồm Môn thể thao, Sản phẩm, Hàng mới, Khuyến mãi, Thương hiệu và Tư vấn. Tìm kiếm luôn hiện diện ở đầu trang; giỏ hàng, yêu thích và tài khoản dùng biểu tượng có nhãn hỗ trợ truy cập.")

add_page_break(doc)

# Page 8 Homepage architecture
add_section_intro(
    doc,
    8,
    "Kiến trúc trang chủ",
    "Trang chủ vừa giới thiệu định vị vừa đưa khách vào danh mục phù hợp. Mỗi khu vực có một vai trò bán hàng cụ thể và tránh lặp lại cùng một lời kêu gọi hành động.",
)
home_sections = [
    ("01", "Header hai tầng", "Tìm kiếm, tài khoản, yêu thích, giỏ hàng và menu danh mục"),
    ("02", "Hero bộ sưu tập", "Thông điệp chính, hai hành động và sản phẩm nổi bật"),
    ("03", "Thanh cam kết", "Giao hàng, đổi trả, chính hãng và thanh toán"),
    ("04", "Sản phẩm mới", "Thẻ sản phẩm có giá, đánh giá, yêu thích và thêm giỏ"),
    ("05", "Flash sale", "Đếm ngược, tiến độ đã bán và ưu đãi có thời hạn"),
    ("06", "Mua theo môn", "Cầu lông, pickleball và tennis bằng hình ảnh rõ chủ đề"),
    ("07", "Bán chạy", "Tab lọc theo môn để so sánh nhanh"),
    ("08", "Chiến dịch", "Hai banner cho giày và phụ kiện"),
    ("09", "Thương hiệu", "Logo chính hãng và đường dẫn đến danh mục"),
    ("10", "Tư vấn và đăng ký", "Bài viết hữu ích, newsletter và footer thông tin"),
]
table = doc.add_table(rows=1, cols=3)
for i, value in enumerate(("Thứ tự", "Khu vực", "Vai trò")):
    table.rows[0].cells[i].text = value
for row in home_sections:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
format_table(table, widths=[1.8, 4.5, 10.0], font_size=8.5)

add_heading(doc, "Nhịp nội dung", 1)
add_body(doc, "Trang xen kẽ khu vực khám phá, sản phẩm và bằng chứng tin cậy. Trên điện thoại, nội dung giữ nguyên thứ tự nhưng sản phẩm chuyển thành hai cột, banner xếp dọc và menu mở bằng ngăn trượt.")

add_page_break(doc)

# Page 9 UX and interactions
add_section_intro(
    doc,
    9,
    "Trải nghiệm và tương tác",
    "Các tương tác mẫu cần đủ thật để khách hàng hiểu cách website vận hành. Mọi trạng thái phải phản hồi ngay, không làm thay đổi bố cục bất ngờ và vẫn dùng được bằng bàn phím.",
)
interactions = [
    ("Tìm kiếm", "Gợi ý sản phẩm và danh mục khi nhập", "Có nhãn, nút xóa và trạng thái không tìm thấy"),
    ("Yêu thích", "Bật tắt tại thẻ và trang sản phẩm", "Thông báo trạng thái đã lưu và giữ lựa chọn trong phiên"),
    ("Thêm giỏ", "Cập nhật số lượng trên biểu tượng giỏ", "Xác nhận ngắn, không chuyển trang bắt buộc"),
    ("Tab lọc", "Lọc sản phẩm bán chạy theo môn", "Tab được chọn rõ và hỗ trợ phím mũi tên"),
    ("Menu di động", "Ngăn trượt chứa nhóm điều hướng", "Khóa cuộn nền, đóng bằng Esc và trả focus đúng vị trí"),
    ("Đếm ngược", "Mô phỏng thời gian flash sale", "Dừng cập nhật khi khu vực không hiển thị"),
    ("Hero slider", "Chuyển bộ sưu tập", "Có nút trước sau, chỉ báo và dừng khi focus hoặc giảm chuyển động"),
    ("Lên đầu trang", "Xuất hiện sau khi cuộn", "Tên truy cập rõ và focus không bị che bởi header"),
]
table = doc.add_table(rows=1, cols=3)
for i, value in enumerate(("Tương tác", "Hành vi", "Yêu cầu trải nghiệm")):
    table.rows[0].cells[i].text = value
for row in interactions:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
format_table(table, widths=[3.2, 5.5, 7.6], font_size=8.45)

add_heading(doc, "Tiêu chí responsive", 1)
add_bullet(doc, "Desktop từ 1280 px: nội dung rộng, danh sách sản phẩm nhiều cột và menu đầy đủ.")
add_bullet(doc, "Tablet từ 768 đến 1279 px: giảm số cột, giữ tìm kiếm và điều hướng ưu tiên.")
add_bullet(doc, "Mobile dưới 768 px: menu hamburger, sản phẩm hai cột và không có cuộn ngang.")

add_page_break(doc)

# Page 10 Visual and content system
add_section_intro(
    doc,
    10,
    "Hệ thống hình ảnh và nội dung",
    "Nhận diện cần tạo cảm giác thể thao cao cấp nhưng không lấn át sản phẩm. Hình ảnh phải giúp khách quan sát món hàng, hiểu bối cảnh sử dụng và nhận ra sự khác biệt giữa các dòng sản phẩm.",
)
visual = [
    ("Màu chính", "Cam #F4511E", "Nút chính, giá khuyến mãi, chỉ báo và chi tiết nhấn"),
    ("Màu chữ", "Đen #171717", "Tiêu đề, nội dung quan trọng và nền footer"),
    ("Nền", "Trắng và #F7F7F7", "Tách khu vực mà không tạo quá nhiều khung"),
    ("Viền", "#E9E9E9", "Phân tách bảng, trường nhập và thẻ sản phẩm"),
    ("Typography", "Be Vietnam Pro hoặc Inter", "Tiếng Việt rõ, số giá dễ quét và cấp bậc nhất quán"),
    ("Chuyển động", "200 đến 300 ms", "Hover, drawer, tab và phản hồi thao tác"),
]
table = doc.add_table(rows=1, cols=3)
for i, value in enumerate(("Thành phần", "Quy ước", "Cách sử dụng")):
    table.rows[0].cells[i].text = value
for row in visual:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
format_table(table, widths=[3.2, 4.2, 9.0], font_size=8.8)

add_heading(doc, "Nguyên tắc hình ảnh", 1)
add_bullet(doc, "Ảnh sản phẩm dùng tỷ lệ 1:1, nền sạch và không cắt mất chi tiết quan trọng.")
add_bullet(doc, "Ảnh danh mục dùng tỷ lệ 4:3, thể hiện đúng môn chơi và bối cảnh sử dụng.")
add_bullet(doc, "Hero dùng tỷ lệ gần 16:5, chừa vùng an toàn cho tiêu đề và nút hành động.")
add_bullet(doc, "Logo thương hiệu dùng tài sản chính thức, giữ đúng tỷ lệ và khoảng trống bảo vệ.")

add_heading(doc, "Giọng nội dung", 1)
add_body(doc, "Nội dung ngắn, cụ thể và có cơ sở từ đặc tính sản phẩm. Tránh lời hứa hiệu suất không được nhà sản xuất xác nhận. Bài tư vấn nên trả lời một nhu cầu rõ như chọn trọng lượng vợt, chọn đế giày hoặc bảo quản dây vợt.")

add_page_break(doc)

# Page 11 Roadmap
add_section_intro(
    doc,
    11,
    "Lộ trình triển khai",
    "Lộ trình chia theo đầu ra có thể xem và kiểm tra. Thời gian là khung tham khảo; phạm vi cuối cùng được xác nhận sau khi có danh mục sản phẩm, quy trình bán hàng và hệ thống vận hành của khách hàng.",
)
roadmap = [
    ("Giai đoạn 1", "Tuần 1 đến 3", "Website mẫu", "Trang chủ hoàn chỉnh, dữ liệu giả lập, responsive và tương tác trình diễn"),
    ("Giai đoạn 2", "Tuần 4 đến 7", "Luồng mua hàng", "Danh mục, chi tiết sản phẩm, giỏ hàng, tài khoản và đặt hàng"),
    ("Giai đoạn 3", "Tuần 8 đến 12", "Vận hành và tăng trưởng", "Thanh toán, kho, thành viên, câu lạc bộ, CRM và phân tích"),
]
table = doc.add_table(rows=1, cols=4)
for i, value in enumerate(("Giai đoạn", "Thời gian", "Mục tiêu", "Đầu ra")):
    table.rows[0].cells[i].text = value
for row in roadmap:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
format_table(table, widths=[2.6, 3.0, 3.4, 7.3], font_size=8.8)

add_heading(doc, "Cách nghiệm thu từng giai đoạn", 1)
add_numbered(doc, "01", "Trình diễn", "Mỗi luồng chính có thể thao tác từ đầu đến cuối trên môi trường thử nghiệm.")
add_numbered(doc, "02", "Thiết bị", "Kiểm tra tối thiểu ở 375 px, 768 px, 1024 px và 1440 px.")
add_numbered(doc, "03", "Chất lượng", "Không có lỗi giao diện nghiêm trọng, lỗi TypeScript hoặc lỗi console ảnh hưởng hành vi.")
add_numbered(doc, "04", "Nội dung", "Tên sản phẩm, giá, chính sách và hình ảnh được khách hàng xác nhận trước khi phát hành.")

add_heading(doc, "Điểm cần khách hàng cung cấp", 1)
add_bullet(doc, "Tên thương hiệu, logo, thông tin liên hệ và chính sách bán hàng.")
add_bullet(doc, "Danh mục sản phẩm, giá, tồn kho, hình ảnh và thông số kỹ thuật.")
add_bullet(doc, "Quy trình giao hàng, đổi trả, bảo hành và thanh toán dự kiến.")

add_page_break(doc)

# Page 12 Data model and next steps
add_section_intro(
    doc,
    12,
    "Mô hình dữ liệu và bước tiếp theo",
    "Các đối tượng dưới đây tạo khung thống nhất giữa nội dung, thiết kế và phát triển. Giai đoạn tài liệu chưa tạo API hoặc kết nối backend; cấu trúc được dùng để chuẩn bị dữ liệu mẫu và tránh thiếu thông tin khi triển khai.",
)
models = [
    ("Product", "Tên, thương hiệu, môn chơi, giá, giá cũ, khuyến mãi, đánh giá, biến thể và tồn kho"),
    ("SportCategory", "Tên môn, hình ảnh, mô tả và nhóm sản phẩm liên quan"),
    ("Promotion", "Tên chiến dịch, thời gian, điều kiện và danh sách sản phẩm áp dụng"),
    ("Article", "Tiêu đề, chủ đề, ảnh, tóm tắt, nội dung và ngày xuất bản"),
    ("Brand", "Tên thương hiệu, logo, mô tả và danh mục sản phẩm"),
]
table = doc.add_table(rows=1, cols=2)
table.rows[0].cells[0].text = "Đối tượng"
table.rows[0].cells[1].text = "Thông tin nghiệp vụ"
for row in models:
    cells = table.add_row().cells
    cells[0].text = row[0]
    cells[1].text = row[1]
format_table(table, widths=[4.0, 12.3], font_size=9.1)

add_heading(doc, "Bước làm việc tiếp theo", 1)
add_numbered(doc, "01", "Duyệt định hướng", "Chọn concept chính và xác nhận nhóm khách hàng ưu tiên.")
add_numbered(doc, "02", "Chuẩn bị nội dung", "Thu thập sản phẩm, chính sách, logo và hình ảnh có quyền sử dụng.")
add_numbered(doc, "03", "Thiết kế giao diện", "Dựng trang chủ và hệ thống thành phần theo hướng đã duyệt.")
add_numbered(doc, "04", "Phát triển bản mẫu", "Xây website responsive với dữ liệu giả lập và các tương tác chính.")

add_heading(doc, "Thông tin xác nhận", 1)
confirm = doc.add_table(rows=3, cols=2)
confirm_data = [
    ("Hướng được chọn", "[PERFORMANCE STORE / SPORT LIFESTYLE / CLUB COMMERCE]"),
    ("Người xác nhận", "[HỌ VÀ TÊN]"),
    ("Ngày xác nhận", "[NGÀY THÁNG NĂM]"),
]
for i, row in enumerate(confirm_data):
    confirm.cell(i, 0).text = row[0]
    confirm.cell(i, 1).text = row[1]
format_table(confirm, header=False, widths=[4.2, 12.1], font_size=9.2)

doc.core_properties.title = "Hồ sơ ý tưởng website kinh doanh dụng cụ thể thao"
doc.core_properties.subject = "Đề xuất sản phẩm, trải nghiệm mua sắm và lộ trình triển khai"
doc.core_properties.author = "[TÊN THƯƠNG HIỆU]"
doc.core_properties.keywords = "website thể thao, thương mại điện tử, hồ sơ ý tưởng"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
